from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
import re
from pathlib import Path
import logging
from processors.module_mapper import ModuleMapper
from excel.excel_manager import ExcelManager

logger = logging.getLogger("CloudAppLogger")

class ModuleNameAdder:
    def __init__(self, workbook_path: Path, sheet_name: str, orig_df):
        self.workbook_path = Path(workbook_path)
        self.sheet_name = sheet_name
        self.orig_df = orig_df

    def process(self):
        logger.info(f"Processing module names for sheet: {self.sheet_name}")
        
        try:
            wb = ExcelManager.load_workbook(self.workbook_path)
            if self.sheet_name not in wb.sheetnames:
                logger.error(f"Sheet {self.sheet_name} not found")
                wb.close()
                return
            
            ws = wb[self.sheet_name]
            
            header_row, slot_cols = self._detect_header_row(ws)
            if header_row is None or not slot_cols:
                logger.error(f"Header row not found in {self.sheet_name}")
                wb.close()
                return
            
            node_headers = self._find_node_headers(ws)
            if not node_headers:
                logger.error(f"No node headers found in {self.sheet_name}")
                wb.close()
                return
            
            module_map = ModuleMapper.build_module_map(self.orig_df, slot_cols, node_headers)
            merged_ranges = list(ws.merged_cells.ranges)
            
            modules_written = self._fill_modules(ws, node_headers, slot_cols, module_map, merged_ranges)
            
            if 'FrontLoading' in self.sheet_name.upper():
                self._write_station_name(ws)
            
            logger.info(f"Filled {modules_written} module names in {self.sheet_name}")
            ExcelManager.save_workbook(wb, self.workbook_path)
            
        except Exception as e:
            logger.error(f"Error processing modules: {e}")
            raise

    def _detect_header_row(self, ws):
        for r in range(1, ws.max_row + 1):
            row_cells = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            row_text = " ".join([str(v) for v in row_cells if v is not None]).upper()
            
            if re.search(r'\bS\s*\d+\b', row_text):
                slot_cols = {}
                for c_idx in range(1, ws.max_column + 1):
                    cv = ws.cell(row=r, column=c_idx).value
                    if cv:
                        m = re.search(r'\bS\s*([1-9]\d*)\b', str(cv))
                        if m:
                            try:
                                slot_cols[int(m.group(1))] = c_idx
                            except Exception:
                                continue
                if slot_cols:
                    return r, slot_cols
        
        return None, {}

    def _find_node_headers(self, ws):
        node_headers = []
        node_re = re.compile(r'(?:SCU)?/?SNU\s*(\d+)', re.IGNORECASE)
        
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                val = ws.cell(row=r, column=c).value
                if val is not None:
                    m = node_re.search(str(val))
                    if m:
                        try:
                            node_num = int(m.group(1))
                            if not any(nh_num == node_num and nh_row == r for nh_num, nh_row in node_headers):
                                node_headers.append((node_num, r))
                        except Exception:
                            continue
        
        return sorted(node_headers, key=lambda x: x[1])

    def _fill_modules(self, ws, node_headers, slot_cols, module_map, merged_ranges):
        modules_written = 0
        
        for nh_node, nh_row in node_headers:
            next_rows = [r for nnum, r in node_headers if r > nh_row]
            zone_end = (min(next_rows) - 1) if next_rows else ws.max_row
            zone_start = nh_row + 1
            
            for slot_p, col_idx in sorted(slot_cols.items()):
                io_mod, redundancy, _ = module_map.get((nh_node, slot_p), ('', False, ''))
                if not io_mod:
                    continue
                
                disp = f"{io_mod} (R)" if redundancy else io_mod
                target_row = zone_start
                target_col = col_idx
                closest_distance = float('inf')
                
                for mr in merged_ranges:
                    try:
                        if mr.min_col <= col_idx <= mr.max_col:
                            if mr.min_row <= zone_end and mr.max_row >= zone_start:
                                distance = abs(mr.min_row - zone_start)
                                if distance < closest_distance:
                                    closest_distance = distance
                                    target_row = mr.min_row
                                    target_col = mr.min_col
                    except Exception:
                        continue
                
                ExcelManager.write_cell_safe(ws, target_row, target_col, disp)
                modules_written += 1
        
        return modules_written

    def _write_station_name(self, ws):
        try:
            station_val = self.orig_df['STATION NAME'].iloc[0] if 'STATION NAME' in self.orig_df.columns else ""
            if station_val:
                for mr in ws.merged_cells.ranges:
                    try:
                        if mr.min_row <= 3 <= mr.max_row and mr.min_col <= 5 and mr.max_col >= 10:
                            ExcelManager.write_cell_safe(ws, 3, 12, station_val)
                            break
                    except Exception:
                        continue
        except Exception as e:
            logger.debug(f"Could not write station name: {e}")
