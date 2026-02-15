"""Test suite for ExcelOutputGenerator"""
import unittest
import pandas as pd
import tempfile
from pathlib import Path
from openpyxl import load_workbook
import sys
import logging

# Add parent directory to path for imports
sys.path.insert(0, str(Path(__file__).parent.parent))

from excel_output import ExcelOutputGenerator


class TestExcelOutputGenerator(unittest.TestCase):
    """Test cases for ExcelOutputGenerator class"""
    
    def setUp(self):
        """Set up test fixtures"""
        self.temp_dir = tempfile.TemporaryDirectory()
        self.temp_path = Path(self.temp_dir.name)
        self.generator = ExcelOutputGenerator(logger=logging.getLogger(__name__))
    
    def tearDown(self):
        """Clean up temporary files"""
        self.temp_dir.cleanup()
    
    def test_write_assignment_to_excel_creates_file(self):
        """Test that write_assignment_to_excel creates a valid Excel file"""
        # Create sample data
        df_assignment = pd.DataFrame({
            'PID_TAG': ['TAG001', 'TAG002', 'TAG003'],
            'Description': ['Signal 1', 'Signal 2', 'Signal 3'],
            'IO_Type': ['AI', 'DI', 'DO'],
            'Module_Name': ['AI-001', 'DI-001', 'DO-001'],
            'Channel': [1, 2, 3],
            'Node': ['Node1', 'Node1', 'Node1'],
            'Slot': ['1', '2', '3'],
            'Signal_Type': ['Input', 'Input', 'Output'],
            'Utilization': ['33.3%', '50%', '25%']
        })
        
        assignment_table = pd.DataFrame({
            'IO_Type': ['AI', 'DI', 'DO'],
            'Module_Name': ['AI-001', 'DI-001', 'DO-001'],
            'Signals': [1, 1, 1],
            'Wired_Spares': [0, 0, 0],
            'Blank_Channels': [15, 15, 15]
        })
        
        output_file = self.temp_path / "test_assignment.xlsx"
        
        # Write Excel file
        result = self.generator.write_assignment_to_excel(
            str(output_file),
            df_assignment,
            assignment_table
        )
        
        # Verify file was created
        self.assertTrue(result)
        self.assertTrue(output_file.exists())
    
    def test_write_assignment_to_excel_sheet_structure(self):
        """Test that sheets have correct structure and headers"""
        df_assignment = pd.DataFrame({
            'PID_TAG': ['TAG001'],
            'Description': ['Signal 1'],
            'IO_Type': ['AI'],
            'Module_Name': ['AI-001'],
            'Channel': [1],
            'Node': ['Node1'],
            'Slot': ['1'],
            'Signal_Type': ['Input'],
            'Utilization': ['33.3%']
        })
        
        assignment_table = pd.DataFrame({
            'IO_Type': ['AI'],
            'Module_Name': ['AI-001'],
            'Signals': [1],
            'Wired_Spares': [0],
            'Blank_Channels': [15]
        })
        
        output_file = self.temp_path / "test_structure.xlsx"
        self.generator.write_assignment_to_excel(str(output_file), df_assignment, assignment_table)
        
        # Verify workbook structure
        wb = load_workbook(str(output_file))
        self.assertIn('Signal_Assignments', wb.sheetnames)
        self.assertIn('Module_Allocation', wb.sheetnames)
        
        # Verify Signal_Assignments sheet
        ws_signals = wb['Signal_Assignments']
        expected_headers = ['PID_TAG', 'Description', 'IO_Type', 'Module_Name', 'Channel',
                          'Node', 'Slot', 'Signal_Type', 'Utilization']
        actual_headers = [cell.value for cell in ws_signals[1]]
        self.assertEqual(actual_headers, expected_headers)
        
        # Verify data row
        self.assertEqual(ws_signals['A2'].value, 'TAG001')
        self.assertEqual(ws_signals['C2'].value, 'AI')
        
        # Verify Module_Allocation sheet
        ws_modules = wb['Module_Allocation']
        expected_module_headers = ['IO_Type', 'Module_Name', 'Signals', 'Wired_Spares', 'Blank_Channels']
        actual_module_headers = [cell.value for cell in ws_modules[1]]
        self.assertEqual(actual_module_headers, expected_module_headers)
        
        # Verify module data
        self.assertEqual(ws_modules['A2'].value, 'AI')
        self.assertEqual(ws_modules['B2'].value, 'AI-001')
        self.assertEqual(ws_modules['C2'].value, 1)
    
    def test_write_assignment_to_excel_header_formatting(self):
        """Test that headers have correct formatting"""
        df_assignment = pd.DataFrame({
            'PID_TAG': ['TAG001'],
            'Description': ['Signal 1'],
            'IO_Type': ['AI'],
            'Module_Name': ['AI-001'],
            'Channel': [1],
            'Node': ['Node1'],
            'Slot': ['1'],
            'Signal_Type': ['Input'],
            'Utilization': ['33.3%']
        })
        
        assignment_table = pd.DataFrame({
            'IO_Type': ['AI'],
            'Module_Name': ['AI-001'],
            'Signals': [1],
            'Wired_Spares': [0],
            'Blank_Channels': [15]
        })
        
        output_file = self.temp_path / "test_formatting.xlsx"
        self.generator.write_assignment_to_excel(str(output_file), df_assignment, assignment_table)
        
        wb = load_workbook(str(output_file))
        ws = wb['Signal_Assignments']
        
        # Check header formatting
        header_cell = ws['A1']
        self.assertTrue(header_cell.font.bold)
        self.assertEqual(header_cell.font.color.rgb, '00FFFFFF')  # White color
        self.assertEqual(header_cell.fill.start_color.rgb, '00366092')  # Blue background
    
    def test_write_assignment_to_excel_empty_dataframe(self):
        """Test handling of empty DataFrames"""
        df_assignment = pd.DataFrame({
            'PID_TAG': [],
            'Description': [],
            'IO_Type': [],
            'Module_Name': [],
            'Channel': [],
            'Node': [],
            'Slot': [],
            'Signal_Type': [],
            'Utilization': []
        })
        
        assignment_table = pd.DataFrame({
            'IO_Type': [],
            'Module_Name': [],
            'Signals': [],
            'Wired_Spares': [],
            'Blank_Channels': []
        })
        
        output_file = self.temp_path / "test_empty.xlsx"
        result = self.generator.write_assignment_to_excel(str(output_file), df_assignment, assignment_table)
        
        # Should still create file with headers
        self.assertTrue(result)
        self.assertTrue(output_file.exists())
        
        wb = load_workbook(str(output_file))
        ws = wb['Signal_Assignments']
        self.assertEqual(ws.max_row, 1)  # Only header row
    
    def test_write_distribution_summary_to_excel_creates_file(self):
        """Test that write_distribution_summary_to_excel creates a valid file"""
        distribution_plan = {
            'AI': {
                'module_name': 'AI-001',
                'signals': 5,
                'wired_spares': 1,
                'blank_channels': 10,
                'total_items': 6,
                'modules_required': 1,
                'channels_per_module': 16,
                'total_available_channels': 16,
                'module_specs': [{
                    'module_instance': 'AI-001_1',
                    'module_index': 1,
                    'signals': 5,
                    'wired_spares': 1,
                    'blank_channels': 10,
                    'total_items': 6,
                    'capacity': 16
                }]
            },
            'DI': {
                'module_name': 'DI-001',
                'signals': 10,
                'wired_spares': 2,
                'blank_channels': 4,
                'total_items': 12,
                'modules_required': 1,
                'channels_per_module': 16,
                'total_available_channels': 16,
                'module_specs': [{
                    'module_instance': 'DI-001_1',
                    'module_index': 1,
                    'signals': 10,
                    'wired_spares': 2,
                    'blank_channels': 4,
                    'total_items': 12,
                    'capacity': 16
                }]
            }
        }
        
        output_file = self.temp_path / "test_distribution.xlsx"
        result = self.generator.write_distribution_summary_to_excel(str(output_file), distribution_plan)
        
        self.assertTrue(result)
        self.assertTrue(output_file.exists())
    
    def test_write_distribution_summary_sheet_structure(self):
        """Test distribution summary sheet structure"""
        distribution_plan = {
            'AI': {
                'module_name': 'AI-001',
                'signals': 5,
                'wired_spares': 1,
                'blank_channels': 10,
                'total_items': 6,
                'modules_required': 1,
                'channels_per_module': 16,
                'total_available_channels': 16,
                'module_specs': [{
                    'module_instance': 'AI-001_1',
                    'module_index': 1,
                    'signals': 5,
                    'wired_spares': 1,
                    'blank_channels': 10,
                    'total_items': 6,
                    'capacity': 16
                }]
            }
        }
        
        output_file = self.temp_path / "test_dist_structure.xlsx"
        self.generator.write_distribution_summary_to_excel(str(output_file), distribution_plan)
        
        wb = load_workbook(str(output_file))
        self.assertIn('Distribution_Summary', wb.sheetnames)
        self.assertIn('Module_Details', wb.sheetnames)
        
        # Check Distribution_Summary sheet
        ws = wb['Distribution_Summary']
        expected_headers = ['IO_Type', 'Module_Name', 'Signals', 'Wired_Spares', 
                          'Blank_Channels', 'Total_Items', 'Modules_Required',
                          'Channels_Per_Module', 'Total_Available', 'Utilization_%']
        actual_headers = [cell.value for cell in ws[1]]
        self.assertEqual(actual_headers, expected_headers)
        
        # Check data
        self.assertEqual(ws['A2'].value, 'AI')
        self.assertEqual(ws['B2'].value, 'AI-001')
        self.assertEqual(ws['C2'].value, 5)
    
    def test_write_distribution_summary_with_mounting_table(self):
        """Test distribution summary with mounting table"""
        distribution_plan = {
            'AI': {
                'module_name': 'AI-001',
                'signals': 5,
                'wired_spares': 1,
                'blank_channels': 10,
                'total_items': 6,
                'modules_required': 1,
                'channels_per_module': 16,
                'total_available_channels': 16,
                'module_specs': [{
                    'module_instance': 'AI-001_1',
                    'module_index': 1,
                    'signals': 5,
                    'wired_spares': 1,
                    'blank_channels': 10,
                    'total_items': 6,
                    'capacity': 16
                }]
            }
        }
        
        mount_table = pd.DataFrame({
            'Slot_1': ['AI-001_1'],
            'Slot_2': [None],
            'Slot_3': [None],
            'Slot_4': [None],
            'Slot_5': [None],
            'Slot_6': [None],
            'Slot_7': [None],
            'Slot_8': [None],
            'Slot_9': [None],
            'Slot_10': [None],
            'Slot_11': [None],
            'Slot_12': [None]
        }, index=['Node_1'])
        
        output_file = self.temp_path / "test_dist_with_mount.xlsx"
        result = self.generator.write_distribution_summary_to_excel(
            str(output_file),
            distribution_plan,
            mount_table
        )
        
        self.assertTrue(result)
        self.assertTrue(output_file.exists())
        
        wb = load_workbook(str(output_file))
        self.assertIn('Mounting_Table', wb.sheetnames)
        
        ws_mount = wb['Mounting_Table']
        # Verify mounting table structure
        self.assertEqual(ws_mount['A1'].value, 'Slot_1')
        self.assertEqual(ws_mount['A2'].value, 'Node_1')
    
    def test_auto_adjust_columns(self):
        """Test that columns are auto-adjusted"""
        df_assignment = pd.DataFrame({
            'PID_TAG': ['VERY_LONG_TAG_NAME_FOR_TESTING'],
            'Description': ['S'],
            'IO_Type': ['AI'],
            'Module_Name': ['AI-001'],
            'Channel': [1],
            'Node': ['N1'],
            'Slot': ['1'],
            'Signal_Type': ['Input'],
            'Utilization': ['33.3%']
        })
        
        assignment_table = pd.DataFrame({
            'IO_Type': ['AI'],
            'Module_Name': ['AI-001'],
            'Signals': [1],
            'Wired_Spares': [0],
            'Blank_Channels': [15]
        })
        
        output_file = self.temp_path / "test_columns.xlsx"
        self.generator.write_assignment_to_excel(str(output_file), df_assignment, assignment_table)
        
        wb = load_workbook(str(output_file))
        ws = wb['Signal_Assignments']
        
        # Check that column width is adjusted (should be > minimum)
        col_width = ws.column_dimensions['A'].width
        self.assertIsNotNone(col_width)
        self.assertGreater(col_width, 1)


if __name__ == '__main__':
    unittest.main()
