from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
import re
import logging
from excel.excel_manager import ExcelManager

logger = logging.getLogger("CloudAppLogger")

class StationWriter:
    @staticmethod
    def write_station_info(workbook_path, sheet_names, station_val):
        if not station_val:
            logger.info("Station value is empty, skipping")
            return
        
        logger.info(f"Writing station information to {len(sheet_names)} sheets")
        
        try:
            wb = ExcelManager.load_workbook(workbook_path)
            
            for sheet_name in sheet_names:
                if sheet_name not in wb.sheetnames:
                    continue
                
                ws = wb[sheet_name]
                node_headers = StationWriter._find_node_headers(ws)
                
                if not node_headers:
                    continue
                
                for node_num, nh_row in node_headers:
                    next_rows = [r for nnum, r in node_headers if r > nh_row]
                    zone_end = (min(next_rows) - 1) if next_rows else ws.max_row
                    
                    channel_rows = StationWriter._find_channel_rows(ws, nh_row, zone_end)
                    if not channel_rows:
                        continue
                    
                    top_row = channel_rows[0]
                    bottom_row = channel_rows[-1]
                    
                    mr = StationWriter._find_am_an_merge(ws, top_row, bottom_row)
                    target_cell = ws.cell(row=mr.min_row, column=mr.min_col) if mr else ws.cell(row=top_row, column=39)
                    
                    target_cell.value = f"{station_val} // Node : {node_num}"
                    try:
                        target_cell.font = Font(bold=True, color='FF000000')
                        target_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False, text_rotation=90)
                    except Exception:
                        pass
            
            ExcelManager.save_workbook(wb, workbook_path)
            logger.info("Station information written successfully")
            
        except Exception as e:
            logger.error(f"Failed to write station info: {e}")
            raise

    @staticmethod
    def _find_node_headers(ws):
        node_headers = []
        node_re = re.compile(r'(?:SCU|SNU)/?SNU?(\d+)', re.IGNORECASE)
        
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                val = ws.cell(row=r, column=c).value
                if val is not None:
                    m = node_re.search(str(val))
                    if m:
                        try:
                            node_headers.append((int(m.group(1)), r))
                            break
                        except Exception:
                            continue
        
        return sorted(node_headers, key=lambda x: x[1])

    @staticmethod
    def _find_channel_rows(ws, node_row, zone_end):
        channel_rows = []
        for r in range(node_row + 1, zone_end + 1):
            v = ws.cell(row=r, column=1).value
            try:
                if 1 <= int(v) <= 16:
                    channel_rows.append(r)
            except Exception:
                continue
        return channel_rows

    @staticmethod
    def _find_am_an_merge(ws, top_row, bottom_row):
        for mr in ws.merged_cells.ranges:
            try:
                if mr.min_col <= 39 and mr.max_col >= 40:
                    if mr.min_row <= top_row <= mr.max_row or mr.min_row <= bottom_row <= mr.max_row:
                        return mr
            except Exception:
                continue
        return None
