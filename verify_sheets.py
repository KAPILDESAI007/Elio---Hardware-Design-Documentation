from openpyxl import load_workbook

wb = load_workbook('Design Input Review_2025-12-25.xlsx')
print(f'Sheets in workbook: {wb.sheetnames}')
print(f'Assigned sheet rows: {wb["Assigned"].max_row}')
print(f'Unassigned sheet rows: {wb["Unassigned"].max_row}')
if 'Summary' in wb.sheetnames:
    print(f'Summary sheet rows: {wb["Summary"].max_row}')
    # Check first few cells of Summary
    summary_ws = wb['Summary']
    print(f'\nFirst 10 rows of Summary sheet:')
    for i, row in enumerate(summary_ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        print(f'Row {i}: {row}')
