from flask import Flask, render_template, request, send_file, jsonify
from pathlib import Path
import pandas as pd
import re
import time
import shutil
import logging
from datetime import datetime
import sys
import importlib.util
import json

from config import TEMPLATE_PATH, UPLOAD_FOLDER
from logger_config import setup_logger
from processors import PIDTagFiller, ModuleNameAdder, StationWriter, ModuleDetailsFiller

# Load DesignInputReview from file with spaces in name
design_review_path = Path(__file__).parent / "Design Input Review.py"
spec = importlib.util.spec_from_file_location("design_input_review", design_review_path)
design_review_module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(design_review_module)
DesignInputReview = design_review_module.DesignInputReview

app = Flask(__name__, template_folder='templates', static_folder='static')
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024
app.config['UPLOAD_FOLDER'] = str(UPLOAD_FOLDER)

logger = setup_logger()

class LogCapture(logging.Handler):
    def __init__(self):
        super().__init__()
        self.records = []
    
    def emit(self, record):
        if record.levelno >= logging.INFO:
            self.records.append({
                'level': record.levelname,
                'message': record.getMessage(),
                'time': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            })

log_capture = LogCapture()
logger.addHandler(log_capture)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/process', methods=['POST'])
def process():
    """Process for Nest Loading & IO Assignment"""
    try:
        log_capture.records = []
        
        customer_file = request.files.get('customer_file')
        template_file = request.files.get('template_file')
        
        if not customer_file:
            return jsonify({'error': 'Customer file is required'}), 400
        
        customer_path = UPLOAD_FOLDER / f"customer_{int(time.time())}.xls"
        customer_file.save(customer_path)
        
        template_path = TEMPLATE_PATH
        if template_file:
            template_path = UPLOAD_FOLDER / f"template_{int(time.time())}.xlsx"
            template_file.save(template_path)
        
        logger.info(f"Reading customer file")
        df = pd.read_excel(customer_path)
        logger.info(f"Customer file loaded successfully with {len(df)} rows")
        
        cabinet_name = df["SYSTEM_CABINET"].iloc[0] if "SYSTEM_CABINET" in df.columns else ""
        station_name = df["STATION NAME"].iloc[0] if "STATION NAME" in df.columns else ""
        
        invalid_chars = r'[:\\/*?\[\]]'
        safe_cabinet_name = re.sub(invalid_chars, '_', str(cabinet_name).strip())
        
        logger.info(f"Cabinet: {cabinet_name}, Station: {station_name}")
        
        output_file = UPLOAD_FOLDER / f"output_{int(time.time())}.xlsx"
        
        logger.info("Copying template file")
        shutil.copy2(template_path, output_file)
        
        from openpyxl import load_workbook
        wb = load_workbook(output_file, keep_links=True)
        
        if "Front_Template" in wb.sheetnames:
            wb["Front_Template"].title = safe_cabinet_name
            logger.info(f"Renamed Front_Template to {safe_cabinet_name}")
        
        if "FrontLoading" in wb.sheetnames:
            front_loading_name = f"{station_name}_FrontLoading" if station_name else "FrontLoading"
            front_loading_name = front_loading_name[:31]
            wb["FrontLoading"].title = front_loading_name
            logger.info(f"Renamed FrontLoading to {front_loading_name}")
        else:
            front_loading_name = "FrontLoading"
        
        wb.save(output_file)
        wb.close()
        
        logger.info(f"Processing {safe_cabinet_name} sheet")
        pid_filler = PIDTagFiller(output_file, safe_cabinet_name, df)
        pid_filler.run()
        
        time.sleep(0.5)
        
        logger.info(f"Processing {front_loading_name} sheet")
        module_adder = ModuleNameAdder(output_file, front_loading_name, df)
        module_adder.run()
        
        time.sleep(0.5)
        
        logger.info("Filling module details")
        module_details_filler = ModuleDetailsFiller(output_file, safe_cabinet_name, df)
        module_details_filler.run()
        
        time.sleep(0.5)
        
        logger.info("Filling FrontLoading details")
        frontloading_details_filler = ModuleDetailsFiller(output_file, front_loading_name, df)
        frontloading_details_filler.run()
        
        time.sleep(0.5)
        
        logger.info("Writing station information")
        StationWriter.add_station_in_am_an(output_file, safe_cabinet_name, station_name)
        StationWriter.add_station_in_am_an(output_file, front_loading_name, station_name)
        
        logger.info("Processing completed successfully")
        
        return jsonify({
            'success': True,
            'output_file': str(output_file),
            'logs': log_capture.records
        })
        
    except Exception as e:
        logger.exception(f"Processing failed: {str(e)}")
        return jsonify({'error': str(e), 'logs': log_capture.records}), 500

@app.route('/api/check-columns', methods=['POST'])
def check_columns():
    """Check which columns are available in the uploaded file"""
    try:
        input_file = request.files.get('input_file')
        
        if not input_file:
            return jsonify({'error': 'Input file is required'}), 400
        
        # Save temp file to check columns
        temp_path = UPLOAD_FOLDER / f"check_{int(time.time())}.xls"
        input_file.save(temp_path)
        
        # Check columns using DesignInputReview
        reviewer = DesignInputReview()
        column_info = reviewer.check_available_columns(temp_path)
        
        # Clean up temp file
        temp_path.unlink(missing_ok=True)
        
        return jsonify(column_info)
        
    except Exception as e:
        logger.exception(f"Failed to check columns: {str(e)}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/design-input-review', methods=['POST'])
def design_input_review():
    """Process for Design Input Review"""
    try:
        log_capture.records = []
        
        input_file = request.files.get('input_file')
        
        if not input_file:
            return jsonify({'error': 'Input file is required'}), 400
        
        input_path = UPLOAD_FOLDER / f"design_input_{int(time.time())}.xls"
        input_file.save(input_path)
        
        # Get user inputs
        system_type = request.form.get('system_type', '')
        controller_model = request.form.get('controller_model', '')
        explosion_protection = request.form.get('explosion_protection', '')
        temperature_rating = request.form.get('temperature_rating', '')
        redundancy_types = []
        is_types = []
        io_types = []
        wired_spares = None
        
        # Parse JSON arrays from form data
        try:
            import json
            redundancy_str = request.form.get('redundancy_types', '[]')
            redundancy_types = json.loads(redundancy_str) if redundancy_str else []
            
            is_str = request.form.get('is_types', '[]')
            is_types = json.loads(is_str) if is_str else []
            
            io_types_str = request.form.get('io_types', '[]')
            io_types = json.loads(io_types_str) if io_types_str else ['FIO']
        except json.JSONDecodeError:
            logger.warning("Failed to parse checkbox selections")
        
        wired_spares_str = request.form.get('wired_spares', '')
        if wired_spares_str:
            try:
                wired_spares = float(wired_spares_str)
            except ValueError:
                logger.warning(f"Invalid wired_spares value: {wired_spares_str}")
        
        logger.info(f"Starting Design Input Review with file: {input_path}")
        logger.info(f"  System Type: {system_type}")
        logger.info(f"  Controller Model: {controller_model}")
        logger.info(f"  Explosion Protection: {explosion_protection}")
        logger.info(f"  Temperature Rating: {temperature_rating}")
        logger.info(f"  IO Types: {io_types}")
        logger.info(f"  Redundancy Types: {redundancy_types}")
        logger.info(f"  IS Types: {is_types}")
        logger.info(f"  Wired Spares %: {wired_spares}")
        
        # Create reviewer with user inputs
        reviewer = DesignInputReview(
            system_type=system_type if system_type else None,
            controller_model=controller_model if controller_model else None,
            explosion_protection=explosion_protection if explosion_protection else None,
            temperature_rating=temperature_rating if temperature_rating else None,
            redundancy_types=redundancy_types,
            is_types=is_types,
            wired_spares=wired_spares,
            io_types=io_types if io_types else ['FIO']
        )
        
        # Override project dir to use the uploaded file
        reviewer.df_instruments = pd.read_excel(input_path)
        logger.info(f"[FILE READ DEBUG] Input file shape: {reviewer.df_instruments.shape}")
        logger.info(f"[FILE READ DEBUG] Input file columns: {list(reviewer.df_instruments.columns)}")
        logger.info(f"[FILE READ DEBUG] First few rows:")
        logger.info(f"{reviewer.df_instruments.head().to_string()}")
        
        # Execute review steps
        if not reviewer.extract_required_columns():
            raise Exception("Failed to extract required columns")
        
        if not reviewer.apply_user_inputs():
            raise Exception("Failed to apply user inputs")
        
        if not reviewer.sort_by_pid_tag():
            raise Exception("Failed to sort by PID_TAG")
        
        if not reviewer.read_hardware_config():
            raise Exception("Failed to read hardware configuration")
        
        if not reviewer.read_mounting_rule():
            raise Exception("Failed to read mounting rule")
        
        if not reviewer.read_controller_limits():
            raise Exception("Failed to read controller limits")
        
        # DEBUG: Log state before assign_modules
        logger.info(f"[PRE-ASSIGN DEBUG] df_instruments shape: {reviewer.df_instruments.shape if reviewer.df_instruments is not None else 'None'}")
        logger.info(f"[PRE-ASSIGN DEBUG] df_instruments columns: {list(reviewer.df_instruments.columns) if reviewer.df_instruments is not None else 'None'}")
        logger.info(f"[PRE-ASSIGN DEBUG] df_hardware shape: {reviewer.df_hardware.shape if reviewer.df_hardware is not None else 'None'}")
        logger.info(f"[PRE-ASSIGN DEBUG] IO_type_base unique: {list(reviewer.df_instruments['IO_type_base'].unique()) if reviewer.df_instruments is not None and 'IO_type_base' in reviewer.df_instruments.columns else 'None'}")
        
        if not reviewer.assign_modules():
            raise Exception("Failed to assign modules")
        
        if not reviewer.read_fio_config():
            raise Exception("Failed to read FIO configuration")
        
        if not reviewer.assign_nodes_and_controllers():
            raise Exception("Failed to assign nodes and controllers")
        
        if not reviewer.identify_unassigned():
            raise Exception("Failed to identify unassigned")
        
        # NOTE: Wired spares are ALREADY assigned in assign_modules()
        # Do NOT call generate_wired_spares() as it duplicates spares
        # The old generate_wired_spares() method is obsolete
        reviewer.df_wired_spares = pd.DataFrame()  # Empty - spares already in df_assigned
        
        if not reviewer.generate_output_file():
            raise Exception("Failed to generate output file")
        
        # Find the generated output file
        today = datetime.now().strftime("%Y-%m-%d")
        project_dir = Path(r"C:\Working\Others\Python\Cloud App Projects")
        output_file = project_dir / f"Design Input Review_{today}.xlsx"
        
        logger.info(f"Design Input Review completed successfully. Output: {output_file}")
        
        return jsonify({
            'success': True,
            'output_file': str(output_file),
            'logs': log_capture.records
        })
        
    except Exception as e:
        logger.exception(f"Design Input Review failed: {str(e)}")
        return jsonify({'error': str(e), 'logs': log_capture.records}), 500

@app.route('/api/download/<path:filename>')
def download_file(filename):
    try:
        return send_file(filename, as_attachment=True)
    except Exception as e:
        logger.error(f"Download failed: {str(e)}")
        return jsonify({'error': 'File not found'}), 404

if __name__ == '__main__':
    app.run(debug=True, port=5000)