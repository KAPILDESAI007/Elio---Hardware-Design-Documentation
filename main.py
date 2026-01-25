import pandas as pd
import re
import time
import logging
from pathlib import Path
from openpyxl import load_workbook
import shutil

from config import PROJECT_DIR, TEMPLATE_PATH, OUTPUT_FILE, CUSTOMER_FILE
from logger_config import setup_logger
from processors import PIDTagFiller, ModuleNameAdder, StationWriter, ModuleDetailsFiller

logger = setup_logger()

def main():
    try:
        if not CUSTOMER_FILE.exists():
            raise FileNotFoundError(f"Customer file not found: {CUSTOMER_FILE}")
        
        logger.info(f"Reading customer file: {CUSTOMER_FILE}")
        df = pd.read_excel(CUSTOMER_FILE)
        logger.info(f"Customer file loaded successfully with {len(df)} rows")
        
        cabinet_name = df["SYSTEM_CABINET"].iloc[0] if "SYSTEM_CABINET" in df.columns else ""
        station_name = df["STATION NAME"].iloc[0] if "STATION NAME" in df.columns else ""
        
        invalid_chars = r'[:\\/*?\[\]]'
        safe_cabinet_name = re.sub(invalid_chars, '_', str(cabinet_name).strip())
        
        logger.info(f"Cabinet: {cabinet_name}, Station: {station_name}")
        
        if OUTPUT_FILE.exists():
            logger.info("Output file exists, removing old file")
            OUTPUT_FILE.unlink()
        
        logger.info("Copying template file")
        shutil.copy2(TEMPLATE_PATH, OUTPUT_FILE)
        
        wb = load_workbook(OUTPUT_FILE, keep_links=True)
        
        if "Front_Template" in wb.sheetnames:
            wb["Front_Template"].title = safe_cabinet_name
            logger.info(f"Renamed Front_Template to {safe_cabinet_name}")
        
        if "FrontLoading" in wb.sheetnames:
            front_loading_name = f"{station_name}_FrontLoading" if station_name else "FrontLoading"
            front_loading_name = front_loading_name[:31]
            wb["FrontLoading"].title = front_loading_name
            logger.info(f"Renamed FrontLoading to {front_loading_name}")
        else:
            front_loading_name = "FrontLoading"
        
        wb.save(OUTPUT_FILE)
        wb.close()
        
        logger.info(f"Processing {safe_cabinet_name} sheet with PIDTagFiller")
        pid_filler = PIDTagFiller(OUTPUT_FILE, safe_cabinet_name, df)
        pid_filler.run()
        
        time.sleep(1)
        
        logger.info(f"Processing {front_loading_name} sheet with ModuleNameAdder")
        module_adder = ModuleNameAdder(OUTPUT_FILE, front_loading_name, df)
        module_adder.run()
        
        time.sleep(1)
        
        logger.info("Filling module details in column B")
        module_details = ModuleDetailsFiller(OUTPUT_FILE, safe_cabinet_name, df)
        module_details.run()
        
        time.sleep(1)
        
        logger.info("Filling FrontLoading sheet details")
        frontloading_details = ModuleDetailsFiller(OUTPUT_FILE, front_loading_name, df)
        frontloading_details.run()
        
        time.sleep(1)
        
        logger.info("Writing station information to sheets")
        StationWriter.add_station_in_am_an(OUTPUT_FILE, safe_cabinet_name, station_name)
        StationWriter.add_station_in_am_an(OUTPUT_FILE, front_loading_name, station_name)
        
        logger.info(f"Processing completed successfully. Output: {OUTPUT_FILE}")
        
    except Exception as e:
        logger.error(f"Processing failed: {e}")
        raise

if __name__ == "__main__":
    main()
