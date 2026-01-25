from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
import shutil
import logging

logger = logging.getLogger("CloudAppLogger")

class ExcelManager:
    def __init__(self, project_dir: Path = None):
        self.project_dir = Path(project_dir) if project_dir else Path.cwd()

    @staticmethod
    def copy_template(src: Path, dest: Path):
        try:
            shutil.copy2(src, dest)
            logger.debug(f"Template copied: {src} -> {dest}")
        except Exception as e:
            logger.error(f"Failed to copy template: {e}")
            raise

    @staticmethod
    def load_workbook(path: Path, keep_links=True):
        try:
            return load_workbook(path, keep_links=keep_links, data_only=False)
        except Exception as e:
            logger.error(f"Failed to load workbook {path}: {e}")
            raise

    @staticmethod
    def save_workbook(wb, path: Path):
        try:
            wb.save(path)
            logger.debug(f"Workbook saved: {path}")
        except Exception as e:
            logger.error(f"Failed to save workbook: {e}")
            raise
        finally:
            try:
                wb.close()
            except Exception:
                pass

    @staticmethod
    def write_cell_safe(ws, row: int, col: int, value):
        for merged in ws.merged_cells.ranges:
            try:
                min_r, min_c, max_r, max_c = range_boundaries(str(merged))
                if min_r <= row <= max_r and min_c <= col <= max_c:
                    ws.cell(min_r, min_c).value = value
                    return
            except Exception:
                continue
        ws.cell(row, col).value = value

    @staticmethod
    def find_merged_range(ws, row: int, col: int):
        for mr in ws.merged_cells.ranges:
            try:
                min_r, min_c, max_r, max_c = range_boundaries(str(mr))
                if min_r <= row <= max_r and min_c <= col <= max_c:
                    class MergedRange:
                        pass
                    r = MergedRange()
                    r.min_row, r.min_col, r.max_row, r.max_col = min_r, min_c, max_r, max_c
                    return r
            except Exception:
                continue
        return None
